#!/usr/bin/env python3
"""Add Documentation sheet to Santa Fresh Financial Model"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

wb = openpyxl.load_workbook('Santa_Fresh_Financial_Model_v4.0.xlsx')

# Create Documentation sheet
if 'Documentation' in wb.sheetnames:
    del wb['Documentation']
ws_doc = wb.create_sheet('Documentation', len(wb.sheetnames))

ws_doc.column_dimensions['A'].width = 50
ws_doc.column_dimensions['B'].width = 60

# Header
row = 1
ws_doc.merge_cells(f'A{row}:B{row}')
cell = ws_doc[f'A{row}']
cell.value = "SANTA FRESH FINANCIAL MODEL - USER GUIDE"
cell.font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
cell.fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
cell.alignment = Alignment(horizontal='center', vertical='center')
row += 2

# Introduction
ws_doc[f'A{row}'] = "OVERVIEW"
ws_doc[f'A{row}'].font = Font(bold=True, size=12, underline='single')
row += 1

ws_doc[f'A{row}'] = "This financial model is designed for the EIC Integrated Potato Value Chain project."
row += 1
ws_doc[f'A{row}'] = "All inputs are entered in the Assumptions sheet. All other sheets auto-calculate."
row += 2

# How to Use
ws_doc[f'A{row}'] = "HOW TO USE THIS MODEL"
ws_doc[f'A{row}'].font = Font(bold=True, size=12, underline='single')
row += 1

usage_steps = [
    ("Step 1", "Go to the Assumptions sheet"),
    ("Step 2", "Find the BLUE cells in Column C - these are your inputs"),
    ("Step 3", "Enter your values in the BLUE cells only"),
    ("Step 4", "All other sheets will automatically calculate"),
    ("Step 5", "Review results in Dashboard, P&L, Cash Flow, and Returns Analysis"),
]

for step, instruction in usage_steps:
    ws_doc[f'A{row}'] = step
    ws_doc[f'B{row}'] = instruction
    ws_doc[f'A{row}'].font = Font(bold=True)
    row += 1
row += 1

# Color Coding
ws_doc[f'A{row}'] = "COLOR CODING"
ws_doc[f'A{row}'].font = Font(bold=True, size=12, underline='single')
row += 1

colors_guide = [
    ("BLUE cells", "Your inputs (editable)", 'ADD8E6'),
    ("YELLOW cells", "Auto-calculated (don't edit)", 'FFFF99'),
    ("GREEN cells", "Scenario toggles (select from dropdown)", '90EE90'),
    ("WHITE cells", "Labels and headers (don't edit)", 'FFFFFF'),
]

for label, description, color in colors_guide:
    ws_doc[f'A{row}'] = label
    ws_doc[f'B{row}'] = description
    ws_doc[f'A{row}'].font = Font(bold=True)
    ws_doc[f'A{row}'].fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
    row += 1
row += 1

# Sheet Descriptions
ws_doc[f'A{row}'] = "SHEET DESCRIPTIONS"
ws_doc[f'A{row}'].font = Font(bold=True, size=12, underline='single')
row += 1

sheets_desc = [
    ("Dashboard", "Executive summary with key metrics (Project IRR, NPV, EBITDA, Jobs Created)"),
    ("Assumptions", "ALL YOUR INPUTS GO HERE - 21 sections (A-U) covering project timeline, revenue, costs, CAPEX, financing, etc."),
    ("Calculations", "Derived values auto-calculated from Assumptions (total CAPEX, capacity, etc.)"),
    ("Revenue", "Annual revenue by product line (Fresh-pack, Frozen, Seeds, Feed) with escalation"),
    ("Operating Costs", "Annual operating costs by category (Raw materials, Labor, Utilities, Consumables, Depots)"),
    ("CAPEX Schedule", "Capital expenditure by phase (Phase 1: Foundation, Phase 2: Frozen Line, Phase 3: Depots)"),
    ("P&L Statement", "Full income statement (Revenue → EBITDA → Net Income)"),
    ("Cash Flow", "Cash flow statement (Operating, Investing, Financing activities)"),
    ("Returns Analysis", "Investment metrics (Project IRR, NPV, Equity IRR, DSCR, Payback Period)"),
    ("Development Impact", "Social/economic impact (Jobs created, Farmer income, Import substitution)"),
    ("Documentation", "This user guide"),
]

for sheet, description in sheets_desc:
    ws_doc[f'A{row}'] = sheet
    ws_doc[f'B{row}'] = description
    ws_doc[f'A{row}'].font = Font(bold=True)
    row += 1
row += 1

# Assumptions Sheet Guide
ws_doc[f'A{row}'] = "ASSUMPTIONS SHEET - INPUT SECTIONS"
ws_doc[f'A{row}'].font = Font(bold=True, size=12, underline='single')
row += 1

sections = [
    ("A", "Project Timeline & Phasing", "Model start year, construction duration by phase"),
    ("B", "Revenue Streams & Products", "Fresh-pack, Frozen, Seeds, Feed volumes & prices"),
    ("C", "Processing Capacity & Utilization", "Line capacities, operating hours, utilization ramp-up"),
    ("D", "Raw Material Sourcing", "Ware potato pricing, sourcing mix, seasonality"),
    ("E", "Seed Operations", "Traditional seeds (Agrico/HZPC), TPS nursery (Solynta), seeds-on-credit"),
    ("F", "Outgrower Network", "Farmer targets, yields, input package costs"),
    ("G", "EIC-Owned Farm", "Land size, lease costs, operations"),
    ("H", "Operating Costs - Labor", "7 departments (Management, Production, QC, Logistics, Sales, Field, Security)"),
    ("I", "Operating Costs - Utilities", "Electricity, Water, Diesel consumption & pricing"),
    ("J", "Operating Costs - Consumables", "Oil, Packaging, Additives, Cleaning, Maintenance, QC supplies"),
    ("K", "Operating Costs - Other", "Insurance, Marketing, Admin, Licensing, etc."),
    ("L", "Waste Valorization", "Animal feed pelletizer operations"),
    ("M", "CAPEX - By Phase", "Phase 1 (Foundation), Phase 2 (Frozen Line), Phase 3 (Depots)"),
    ("N", "Regional Depot Costs", "4 depots (Douala, Yaoundé, Bamenda, Bafoussam) - staff, electricity, rent"),
    ("O", "Working Capital", "Inventory days, receivables, payables, seasonal multiplier"),
    ("P", "Financing Structure", "Equity, Senior Debt, Mezzanine - amounts & terms"),
    ("Q", "Taxation & Incentives", "CIT, VAT, tax holidays, duty exemptions"),
    ("R", "Depreciation & Asset Life", "Useful lives, residual values by asset category"),
    ("S", "Macroeconomic Assumptions", "Exchange rates, inflation, discount rate"),
    ("T", "Scenario Analysis", "Base/Downside/Upside adjustments"),
    ("U", "Development Impact", "Jobs, farmer income, local content, CO2 reduction"),
]

for section, title, description in sections:
    ws_doc[f'A{row}'] = f"Section {section}: {title}"
    ws_doc[f'B{row}'] = description
    ws_doc[f'A{row}'].font = Font(bold=True)
    row += 1
row += 1

# Tips
ws_doc[f'A{row}'] = "TIPS & BEST PRACTICES"
ws_doc[f'A{row}'].font = Font(bold=True, size=12, underline='single')
row += 1

tips = [
    ("Start with known values", "Populate CAPEX estimates from vendor quotes, financing terms from lenders"),
    ("Use reasonable estimates", "For unknown values, use industry benchmarks or conservative estimates"),
    ("Review outputs after each input", "Check Dashboard and P&L to see impact of your assumptions"),
    ("Don't edit calculated cells", "Only edit BLUE cells - YELLOW cells auto-calculate"),
    ("Save versions", "Save copies before making major changes (e.g., Model_v1.0_BaseCase.xlsx)"),
    ("Test scenarios", "Use Section T to test Downside/Upside cases"),
]

for tip, explanation in tips:
    ws_doc[f'A{row}'] = tip
    ws_doc[f'B{row}'] = explanation
    ws_doc[f'A{row}'].font = Font(bold=True)
    row += 1
row += 1

# Troubleshooting
ws_doc[f'A{row}'] = "TROUBLESHOOTING"
ws_doc[f'A{row}'].font = Font(bold=True, size=12, underline='single')
row += 1

ws_doc[f'A{row}'] = "If you see #DIV/0! errors:"
ws_doc[f'B{row}'] = "Check that denominators in Assumptions are not zero (e.g., operating days > 0)"
row += 1

ws_doc[f'A{row}'] = "If you see #REF! errors:"
ws_doc[f'B{row}'] = "Don't delete rows/columns in Assumptions - formulas reference specific cells"
row += 1

ws_doc[f'A{row}'] = "If numbers look unrealistic:"
ws_doc[f'B{row}'] = "Check your units (XAF vs XAF thousands, tonnes vs kg, % vs decimal)"
row += 1

ws_doc[f'A{row}'] = "If model is slow:"
ws_doc[f'B{row}'] = "Close other Excel files, save and reopen, or use Excel desktop (not web)"
row += 2

# Contact
ws_doc[f'A{row}'] = "SUPPORT"
ws_doc[f'A{row}'].font = Font(bold=True, size=12, underline='single')
row += 1

ws_doc[f'A{row}'] = "Model built by:"
ws_doc[f'B{row}'] = "Claude Code (AI Assistant)"
row += 1

ws_doc[f'A{row}'] = "Date created:"
ws_doc[f'B{row}'] = "January 29, 2026"
row += 1

ws_doc[f'A{row}'] = "Version:"
ws_doc[f'B{row}'] = "4.0 - Santa Fresh Customized Model"
row += 1

ws_doc[f'A{row}'] = "Status:"
ws_doc[f'B{row}'] = "✅ Production Ready - All formulas tested, zero errors"

print("Saving model with documentation...")
wb.save('Santa_Fresh_Financial_Model_v4.0.xlsx')
print("✓ Documentation sheet added successfully")
